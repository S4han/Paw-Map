"""식약처 '반려동물 동반 가능 업소 현황' 파일 어댑터 — 아키텍처 §6.2, api-analysis §1A.

실측 열: 연번, 업소명, 업종, 지역, 업소주소. 좌표·고유 ID·날짜·동반 조건이 없다.
- 등재 = 동반 가능(OFFICIAL). 미등재 업소에 대한 판정은 만들지 않는다.
- 날짜가 없으므로 checkedAt 을 채우지 않는다(lastVerifiedAt 도 비어 있게 됨).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from ..classify import mfds_category
from ..geocode import Geocoder
from ..models import Evidence, PetConditions, PetPlace
from ..normalize import clean_address, norm_address, norm_name, stable_key
from ..rules import decide_pet_status

CLAIM = "식약처 반려동물 동반출입 음식점 등록"
# 지역 필터: 법정동 시도 코드 → 주소 앞부분
SIDO_PREFIX = {"50": ("제주",), "11": ("서울",), "51": ("강원",), "41": ("경기",)}


@dataclass
class MfdsRow:
    name: str
    business_type: str
    address: str


@dataclass
class MfdsStats:
    rows: int = 0
    in_region: int = 0
    geocoded_address: int = 0
    geocoded_keyword: int = 0
    failed: list[str] = field(default_factory=list)
    duplicates_in_file: int = 0


def read_xlsx(path: Path) -> list[list[str]]:
    import openpyxl  # 가공 환경 전용 의존성

    # 이 파일은 시트 크기 정보가 잘못 기록돼 read_only 로 읽으면 첫 칸만 나온다(STEP 2 실측)
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = max(wb.worksheets, key=lambda ws: ws.max_row)
    return [["" if c is None else str(c).strip() for c in row] for row in sheet.iter_rows(values_only=True)]


def parse_rows(rows: list[list[str]]) -> list[MfdsRow]:
    header_i = next(i for i, r in enumerate(rows[:15]) if "업소명" in r)
    header = rows[header_i]
    col = {name: header.index(name) for name in ("업소명", "업종", "업소주소")}
    out = []
    for r in rows[header_i + 1 :]:
        if len(r) <= max(col.values()) or not r[col["업소명"]]:
            continue
        out.append(MfdsRow(r[col["업소명"]], r[col["업종"]], r[col["업소주소"]]))
    return out


def places(rows: list[MfdsRow], region_code: str, geocoder: Geocoder, fetched_at: str, stats: MfdsStats) -> list[PetPlace]:
    prefixes = SIDO_PREFIX.get(region_code)
    if not prefixes:
        raise ValueError(f"식약처 지역 필터가 정의되지 않은 지역 코드: {region_code}")
    stats.rows = len(rows)
    seen: set[str] = set()
    result: list[PetPlace] = []
    for row in rows:
        if not row.address.startswith(prefixes):
            continue
        stats.in_region += 1
        key = stable_key(norm_name(row.name), norm_address(row.address))
        if key in seen:
            stats.duplicates_in_file += 1
            continue
        seen.add(key)
        address = clean_address(row.address)
        geo = geocoder.geocode(address, row.name)
        if geo is None:
            stats.failed.append(f"{row.name} | {row.address}")
            continue
        if geo.method == "address":
            stats.geocoded_address += 1
        else:
            stats.geocoded_keyword += 1
        evidence = [Evidence(source="MFDS_PET", source_id=key, type="OFFICIAL", pet_status="ALLOWED", fetched_at=fetched_at, claim=CLAIM)]
        status, conflict = decide_pet_status(evidence)
        result.append(
            PetPlace(
                id=f"mfds:{key}",
                name=row.name,
                category=mfds_category(row.business_type, row.name),
                latitude=geo.lat,
                longitude=geo.lng,
                address=address,
                pet_status=status,
                pet_conditions=PetConditions(notes=[f"업종: {row.business_type}"]),
                evidence=evidence,
                conflicting_evidence=conflict,
            )
        )
    return result
